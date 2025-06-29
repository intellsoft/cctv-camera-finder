import tkinter as tk
from tkinter import ttk, messagebox, filedialog, Menu
import subprocess
import socket
import threading
import ipaddress
from ping3 import ping
import concurrent.futures
import queue
from onvif import ONVIFCamera
import os
import json
from cryptography.fernet import Fernet
import openpyxl
import requests
from PIL import Image, ImageTk
import io
from urllib.parse import urlparse
import webbrowser
import platform
from zeroconf import ServiceBrowser, Zeroconf, ServiceListener
from getmac import get_mac_address
import re
import cv2

# --- New Imports for Updates ---
from ttkthemes import ThemedTk

# Global variable for the fixed encryption key
FIXED_ENCRYPTION_KEY = b'v8VZkq4W9E7RjTzOYKbLwSdXyGfChN3s1Mn2PpI0Qt6='

class BonjourListener(ServiceListener):
    """Listener for Bonjour/mDNS service discovery"""
    def __init__(self, app):
        self.app = app
        
    def add_service(self, zc, type_, name):
        info = zc.get_service_info(type_, name)
        if info:
            addresses = info.parsed_addresses()
            if addresses:
                ip = addresses[0]
                port = info.port
                self.app.add_bonjour_device(ip, port)
                
    def remove_service(self, zc, type_, name):
        pass
        
    def update_service(self, zc, type_, name):
        pass

class CmdToolExecutor:
    """
    A helper class to execute CMD commands in a separate thread
    and update a Tkinter Text widget.
    """
    def __init__(self, output_widget, root_after_callback):
        self.output_widget = output_widget
        self.root_after = root_after_callback

    def execute_command(self, command):
        self.root_after(0, lambda: self._update_output_gui(f"Executing: {command}\n\n", clear_previous=True))
        threading.Thread(target=self._run_command_in_thread, args=(command,), daemon=True).start()

    def _run_command_in_thread(self, command):
        try:
            creation_flags = subprocess.CREATE_NO_WINDOW if platform.system() == "Windows" else 0
            process = subprocess.run(
                command,
                shell=True,
                capture_output=True,
                text=True,
                check=False,
                encoding='utf-8',
                errors='replace',
                creationflags=creation_flags,
                timeout=60
            )
            output = process.stdout
            error_output = process.stderr
            self.root_after(0, self._update_output_gui, output, error_output)
        except subprocess.TimeoutExpired:
            self.root_after(0, self._update_output_gui, f"Command '{command}' timed out after 60 seconds.", "", is_error=True)
        except Exception as e:
            self.root_after(0, self._update_output_gui, f"Error executing command: {e}", "", is_error=True)

    def _update_output_gui(self, output, error_output="", clear_previous=False, is_error=False):
        self.output_widget.config(state=tk.NORMAL)
        if clear_previous:
            self.output_widget.delete(1.0, tk.END)
        
        self.output_widget.insert(tk.END, output)
        if error_output:
            self.output_widget.insert(tk.END, f"\n\n--- ERROR ---\n{error_output}")
        if is_error:
            self.output_widget.tag_configure("error", foreground="red")
            self.output_widget.insert(tk.END, "\n\n--- COMMAND FAILED ---\n", "error")

        self.output_widget.see(tk.END)
        self.output_widget.config(state=tk.DISABLED)

class RealTimeNetworkScanner:
    def __init__(self, root):
        self.root = root
        self.root.title("Advanced Network Scanner")
        self.root.state('zoomed')

        self.stop_flag = False
        self.search_thread = None
        self.scanning_active = False
        self.found_count = 0
        
        # تبلیغات چرخشی
        self.ads = [
            {"text": "نرم افزار تبدیل عکس به فیلم تایم لپس", "url": "https://intellsoft.ir/product/time-lapse-photo-to-video/"},
            {"text": "نرم افزار تبدیل فیلم دوربین مداربسته تایم لپس", "url": "https://intellsoft.ir/product/time-lapse-software-with-cctv-playback-film/"},
            {"text": "نرم افزار تایم لپس دوربین مداربسته", "url": "https://intellsoft.ir/product/timelapse-camera-recorder/"}
        ]
        self.current_ad_index = 0
        
        # Bonjour variables
        self.zeroconf = None
        self.bonjour_listener = None
        
        # Floating message setup
        self.floating_message = ttk.Label(
            root, 
            text="", 
            background="#FFFFCC", 
            relief="solid", 
            borderwidth=1,
            padding=(5, 2, 5, 2),
            font=("Arial", 9)
        )
        self.floating_message.place(relx=1.0, rely=1.0, anchor='se', x=-10, y=-10)
        self.floating_message.place_forget()
        
        # --- Worker Queues ---
        self.onvif_queue = queue.Queue()
        self.onvif_threads = []
        self.port_scan_queue = queue.Queue()
        self.port_scan_threads = []

        # --- Encryption Setup ---
        try:
            self.cipher_suite = Fernet(FIXED_ENCRYPTION_KEY)
        except Exception as e:
            messagebox.showerror("Error", f"Encryption initialization error: {e}\nPlease ensure a valid encryption key is set.")
            self.root.destroy()
            return

        # --- Credential Management ---
        self.credential_sets = {} 
        self.camera_ip_to_cred_set = {} 

        # Cache for images
        self.thumbnail_cache = {} 
        self.hover_image_cache = {}

        # Create a small static image for the Treeview column
        self.camera_icon = None
        try:
            img = Image.new('RGB', (16, 16), color = 'green')
            self.camera_icon = ImageTk.PhotoImage(img)
        except Exception as e:
            print(f"Could not create camera icon: {e}")
            self.camera_icon = None

        # Hover image related variables (FIXED)
        self.hover_window = None
        self.hover_image_label = None
        self.current_hover_item = None
        self.current_hover_image_tk = None
        self.hover_image_thread = None
        self.last_hover_thread_id = 0

        # --- Scan Range & Settings Management ---
        self.scan_range = {"start_ip": "", "end_ip": ""}
        self.app_theme = "arc" # Default theme
        self.load_settings()
        self.load_camera_data()

        # --- Set Initial Theme ---
        try:
            self.root.set_theme(self.app_theme)
        except Exception:
            self.app_theme = "arc" # Fallback to a default theme
            self.root.set_theme(self.app_theme)

        
        self.create_widgets()
        self.setup_style()
        self.start_worker_threads()
        
        # Start ad rotation
        self.update_advertisement()

    def update_advertisement(self):
        """Updates the advertisement every 5 seconds"""
        if not hasattr(self, 'ad_label'):
            return
            
        ad = self.ads[self.current_ad_index]
        self.ad_label.config(text=ad["text"])
        
        # Update index for next ad
        self.current_ad_index = (self.current_ad_index + 1) % len(self.ads)
        
        # Schedule next update
        self.root.after(5000, self.update_advertisement)

    def center_window(self, window):
        """Centers a window on the screen"""
        window.update_idletasks()
        width = window.winfo_width()
        height = window.winfo_height()
        x = (window.winfo_screenwidth() // 2) - (width // 2)
        y = (window.winfo_screenheight() // 2) - (height // 2)
        window.geometry(f'+{x}+{y}')

    def show_floating_message(self, message):
        """Shows a temporary message in the bottom-right corner"""
        self.floating_message.config(text=message)
        self.floating_message.place(relx=1.0, rely=1.0, anchor='se', x=-10, y=-10)
        self.floating_message.lift()
        self.root.after(3000, self.hide_floating_message)
    
    def hide_floating_message(self):
        """Hides the floating message"""
        self.floating_message.place_forget()

    def encrypt_password(self, password):
        """Encrypts the password."""
        if not password:
            return ""
        return self.cipher_suite.encrypt(password.encode()).decode()

    def decrypt_password(self, encrypted_password):
        """Decrypts the password."""
        if not encrypted_password:
            return ""
        try:
            return self.cipher_suite.decrypt(encrypted_password.encode()).decode()
        except Exception as e:
            print(f"Decryption error: {e}")
            return "Decryption Failed"

    def load_camera_data(self):
        """Loads camera information and credential sets from file."""
        self.credential_sets = {
            "Default Admin": {"username": "admin", "password": self.encrypt_password("admin")}
        }
        self.camera_ip_to_cred_set = {}

        if os.path.exists('camera_data.json'):
            try:
                with open('camera_data.json', 'r') as f:
                    data = json.load(f)
                    if 'credential_sets' in data:
                        for name, creds in data['credential_sets'].items():
                            self.credential_sets[name] = {
                                'username': creds['username'],
                                'password': creds['password']
                            }
                    if 'camera_ip_to_cred_set' in data:
                        self.camera_ip_to_cred_set = data['camera_ip_to_cred_set']
            except json.JSONDecodeError as e:
                self.show_floating_message(f"Error reading camera_data.json: {e}")
            except Exception as e:
                self.show_floating_message(f"Error loading camera data: {e}")
        
    def save_camera_data(self):
        """Saves camera information and credential sets to file."""
        data = {
            'credential_sets': self.credential_sets,
            'camera_ip_to_cred_set': self.camera_ip_to_cred_set
        }
        try:
            with open('camera_data.json', 'w') as f:
                json.dump(data, f, indent=4)
        except Exception as e:
            self.show_floating_message(f"Error saving camera data: {e}")

    def load_settings(self):
        """Loads application settings, including scan range and theme, from file."""
        if os.path.exists('app_settings.json'):
            try:
                with open('app_settings.json', 'r') as f:
                    settings = json.load(f)
                    self.scan_range = settings.get('scan_range', {"start_ip": "", "end_ip": ""})
                    self.app_theme = settings.get('theme', 'arc') # Load theme
            except (json.JSONDecodeError, KeyError) as e:
                self.show_floating_message(f"Error reading app_settings.json: {e}")
            except Exception as e:
                self.show_floating_message(f"Error loading application settings: {e}")

    def save_settings(self):
        """Saves application settings, including scan range and theme, to file."""
        settings = {
            'scan_range': self.scan_range,
            'theme': self.app_theme # Save theme
        }
        try:
            with open('app_settings.json', 'w') as f:
                json.dump(settings, f, indent=4)
        except Exception as e:
            self.show_floating_message(f"Error saving application settings: {e}")

    def start_worker_threads(self):
        """Starts all worker threads (ONVIF and Port Scan)."""
        for _ in range(10):
            t = threading.Thread(target=self.onvif_worker, daemon=True)
            t.start()
            self.onvif_threads.append(t)
        for _ in range(20):
            t = threading.Thread(target=self.port_scan_worker, daemon=True)
            t.start()
            self.port_scan_threads.append(t)

    def onvif_worker(self):
        """ONVIF discovery worker."""
        while True:
            ip, item_id = self.onvif_queue.get()
            # GEMINI FIX: Removed self.tree.exists(item_id) check from this background thread.
            # The check is correctly performed in update_camera_info which runs in the main thread.
            if self.stop_flag:
                self.onvif_queue.task_done()
                continue
                
            self.root.after(0, self.update_status, f"Discovering ONVIF for {ip}")
            
            cred_set_name = self.camera_ip_to_cred_set.get(ip, "Default Admin")
            cred_info = self.credential_sets.get(cred_set_name)

            username = "admin"
            password = "admin"

            if cred_info:
                username = cred_info['username']
                password = self.decrypt_password(cred_info['password'])

            camera_status, rtsp_url, snapshot_uri = self.check_onvif_camera(ip, username, password)
            
            if camera_status == "Auth Failed" and cred_set_name != "Default Admin":
                default_cred_info = self.credential_sets.get("Default Admin")
                if default_cred_info:
                    default_username = default_cred_info['username']
                    default_password = self.decrypt_password(default_cred_info['password'])
                    camera_status, rtsp_url, snapshot_uri = self.check_onvif_camera(ip, default_username, default_password)
                    if camera_status == "ONVIF Found":
                        self.camera_ip_to_cred_set[ip] = "Default Admin"
                        self.save_camera_data()
                        cred_set_name = "Default Admin"
            
            self.root.after(0, self.update_camera_info, item_id, cred_set_name, rtsp_url, camera_status, snapshot_uri)
            self.onvif_queue.task_done()

    def port_scan_worker(self):
        """Port scanning worker."""
        while True:
            ip, item_id = self.port_scan_queue.get()
            # GEMINI FIX: Removed self.tree.exists(item_id) check from this background thread.
            # The check is correctly performed in update_open_ports which runs in the main thread.
            if self.stop_flag:
                self.port_scan_queue.task_done()
                continue
            
            common_ports = [80, 443, 554, 8080, 8000, 22, 23, 37777, 8899, 81]
            open_ports = []
            
            with concurrent.futures.ThreadPoolExecutor(max_workers=len(common_ports)) as executor:
                future_to_port = {executor.submit(self.scan_single_port, ip, port): port for port in common_ports}
                for future in concurrent.futures.as_completed(future_to_port):
                    port = future_to_port[future]
                    try:
                        if future.result():
                            open_ports.append(port)
                    except Exception:
                        pass
            
            open_ports.sort()
            self.root.after(0, self.update_open_ports, item_id, open_ports)
            self.port_scan_queue.task_done()

    def scan_single_port(self, ip, port):
        """Scans a single port on a given IP."""
        try:
            with socket.socket(socket.AF_INET, socket.SOCK_STREAM) as s:
                s.settimeout(0.5)
                s.connect((ip, port))
                return True
        except (socket.timeout, ConnectionRefusedError):
            return False
        except Exception:
            return False

    def check_onvif_camera(self, ip, username, password):
        """
        Checks for ONVIF camera presence and extracts RTSP and Snapshot URIs.
        """
        rtsp_url = ""
        snapshot_uri = ""
        try:
            # خط اصلاح شده
            cam = ONVIFCamera(ip, 80, username, password)
            media_service = cam.create_media_service()
            profiles = media_service.GetProfiles()
            if not profiles:
                return "No Profiles", "", ""
            
            token = profiles[0].token
            stream_uri = media_service.GetStreamUri({
                'StreamSetup': {'Stream': 'RTP-Unicast', 'Transport': {'Protocol': 'RTSP'}},
                'ProfileToken': token
            })
            rtsp_url = stream_uri.Uri

            try:
                snapshot_uri_response = media_service.GetSnapshotUri({'ProfileToken': token})
                snapshot_uri = snapshot_uri_response.Uri
            except Exception as e:
                print(f"Error getting Snapshot URI for {ip}: {e}")
                snapshot_uri = ""

            return "ONVIF Found", rtsp_url, snapshot_uri
        except Exception as e:
            error_str = str(e).lower()
            if "unauthorized" in error_str or "401" in error_str:
                return "Auth Failed", "", ""
            elif "timed out" in error_str or "timeout" in error_str:
                return "Timeout", "", ""
            else:
                return "Error", "", str(e)

    def setup_style(self):
        self.style = ttk.Style()
        self.style.configure('red.TFrame', background='#ffcccc')
        self.style.configure('yellow.TFrame', background='#ffffcc')
        self.style.configure('purple.TFrame', background='#ffccff')
        self.style.configure('camera_found.TLabel', background='#ccffcc')
        self.style.configure('camera_error.TLabel', background='#ffcccc')
        self.style.configure('bonjour_device.TLabel', background='#cce5ff')
        self.style.configure('mac_dup.TLabel', background='#ffcc99')
        self.style.configure('advertisement.TLabel', background='#e6f7ff', foreground='#0066cc', font=('Tahoma', 10, 'bold'))


    def create_widgets(self):
        self.create_main_menu()
        
        main_frame = ttk.Frame(self.root)
        main_frame.pack(fill=tk.BOTH, expand=True, padx=10, pady=10)

        # Control Panel
        control_frame = ttk.Frame(main_frame)
        control_frame.pack(fill=tk.X, pady=5)

        self.search_btn = ttk.Button(
            control_frame,
            text="Search Network",
            command=self.start_search
        )
        self.search_btn.pack(side=tk.LEFT, padx=5)

        self.stop_btn = ttk.Button(
            control_frame,
            text="Stop Search",
            command=self.stop_search,
            state=tk.DISABLED
        )
        self.stop_btn.pack(side=tk.LEFT, padx=5)
        
        self.bonjour_btn = ttk.Button(
            control_frame,
            text="Scan Bonjour",
            command=self.start_bonjour_scan
        )
        self.bonjour_btn.pack(side=tk.LEFT, padx=5)

        self.manage_creds_btn = ttk.Button(
            control_frame,
            text="Manage Credentials",
            command=self.open_credential_manager
        )
        self.manage_creds_btn.pack(side=tk.LEFT, padx=5)

        self.manage_scan_range_btn = ttk.Button(
            control_frame,
            text="Manage Scan Range",
            command=self.open_scan_range_manager
        )
        self.manage_scan_range_btn.pack(side=tk.LEFT, padx=5)

        # Camera Credentials Selection Frame
        cred_selection_frame = ttk.Frame(control_frame)
        cred_selection_frame.pack(side=tk.RIGHT, padx=10)
        
        ttk.Label(cred_selection_frame, text="Credential Set:").pack(side=tk.LEFT)
        self.cred_set_var = tk.StringVar()
        self.cred_set_combobox = ttk.Combobox(
            cred_selection_frame,
            textvariable=self.cred_set_var,
            state="readonly",
            width=15
        )
        self.cred_set_combobox.pack(side=tk.LEFT, padx=5)
        self.cred_set_combobox.bind("<<ComboboxSelected>>", self.on_cred_set_selected)
        self.update_cred_combobox_options()
        
        self.apply_cred_btn = ttk.Button(
            cred_selection_frame,
            text="Apply to Selected IP",
            command=self.apply_camera_credentials
        )
        self.apply_cred_btn.pack(side=tk.LEFT, padx=5)

        # CMD Tools Menubutton (FIXED: Added dropdown menu)
        self.cmd_tools_menubutton = ttk.Menubutton(
            control_frame,
            text="CMD Tools"
        )
        self.cmd_tools_menubutton.pack(side=tk.LEFT, padx=5)
        
        # Create the dropdown menu
        self.cmd_tools_menu = Menu(self.cmd_tools_menubutton, tearoff=0)
        self.cmd_tools_menubutton["menu"] = self.cmd_tools_menu
        
        # Add categories and commands to the menu
        self._create_cmd_tool_menus()

        # Progress Section
        self.progress_frame = ttk.Frame(main_frame)
        self.progress_frame.pack(fill=tk.X, pady=5)
        
        search_frame = ttk.Frame(self.progress_frame)
        search_frame.pack(side=tk.LEFT, fill=tk.X, expand=True, padx=5)
        ttk.Label(search_frame, text="Search:").pack(side=tk.LEFT)
        self.search_var = tk.StringVar()
        self.search_entry = ttk.Entry(search_frame, textvariable=self.search_var)
        self.search_entry.pack(side=tk.LEFT, fill=tk.X, expand=True)
        self.search_entry.bind("<KeyRelease>", self.filter_results)

        self.progress_label = ttk.Label(
            self.progress_frame,
            text="آماده برای اسکن..."
        )
        self.progress_label.pack(side=tk.LEFT)

        self.progress_bar = ttk.Progressbar(
            self.progress_frame,
            orient=tk.HORIZONTAL,
            mode='determinate'
        )
        self.progress_bar.pack(side=tk.RIGHT, fill=tk.X, expand=True)
        self.progress_bar.pack_forget()

        # Results Treeview
        tree_frame = ttk.Frame(main_frame)
        tree_frame.pack(fill=tk.BOTH, expand=True)

        columns = ('IP', 'MAC', 'Credential Set', 'Open Ports', 'RTSP URL', 'Camera Status', 'Snapshot')
        self.tree = ttk.Treeview(tree_frame, columns=columns, show='headings')
        
        col_widths = {
            'IP': 120, 
            'MAC': 130, 
            'Credential Set': 120,
            'Open Ports': 120,
            'RTSP URL': 200,
            'Camera Status': 120,
            'Snapshot': 80
        }

        for col in columns:
            self.tree.heading(col, text=col, command=lambda c=col: self.sort_treeview_column(c, False))
            self.tree.column(col, width=col_widths.get(col, 100))

        scrollbar = ttk.Scrollbar(tree_frame, orient=tk.VERTICAL, command=self.tree.yview)
        self.tree.configure(yscroll=scrollbar.set)
        scrollbar.pack(side=tk.RIGHT, fill=tk.Y)
        self.tree.pack(fill=tk.BOTH, expand=True)

        self.tree.tag_configure('ip_dup', background='#ffcccc')
        self.tree.tag_configure('camera_found', background='#ccffcc')
        self.tree.tag_configure('camera_error', background='#ffcccc')
        self.tree.tag_configure('bonjour_device', background='#cce5ff')
        self.tree.tag_configure('mac_dup', background='#ffcc99')
        
        self.tree.bind('<<TreeviewSelect>>', self.on_tree_select)
        self.tree.bind("<Button-3>", self.on_right_click)
        self.tree.bind("<Motion>", self._on_tree_motion)
        self.tree.bind("<Leave>", self._on_tree_leave)

        # تبلیغات چرخشی
        ad_frame = ttk.Frame(self.root, relief=tk.RAISED, borderwidth=1)
        ad_frame.pack(side=tk.BOTTOM, fill=tk.X, pady=(5, 0))
        
        # Add a label for the advertisement
        self.ad_label = ttk.Label(
            ad_frame, 
            text="", 
            style='advertisement.TLabel',
            cursor="hand2",
            anchor=tk.CENTER,
            padding=5
        )
        self.ad_label.pack(fill=tk.X, expand=True)
        self.ad_label.bind("<Button-1>", self.on_ad_click)
        
        # Footer Section
        footer_frame = ttk.Frame(self.root)
        footer_frame.pack(side=tk.BOTTOM, fill=tk.X, pady=(0, 5))

        footer_text_part1 = "برنامه نویس : محمد علی عباسپور    "
        website_text = "intellsoft.ir"
        footer_text_part2 = " (نرم افزار تخصصی دوربین مداربسته)"
        
        self.footer_text_widget = tk.Text(
            footer_frame, 
            height=1, 
            wrap="word", 
            font=("Arial", 9),
            bg=self.root.cget('bg'),
            relief="flat",
            highlightthickness=0,
            cursor="arrow"
        )
        self.footer_text_widget.pack(side=tk.LEFT, padx=5)
        
        self.footer_text_widget.insert(tk.END, footer_text_part1)
        self.footer_text_widget.insert(tk.END, website_text, "link")
        self.footer_text_widget.insert(tk.END, footer_text_part2)

        self.footer_text_widget.tag_config("link", foreground="blue", underline=1)
        self.footer_text_widget.tag_bind("link", "<Button-1>", self._open_website_link)
        self.footer_text_widget.tag_bind("link", "<Enter>", self._on_link_enter)
        self.footer_text_widget.tag_bind("link", "<Leave>", self._on_link_leave)
        self.footer_text_widget.config(state="disabled")
        
    def on_ad_click(self, event):
        """Handles click on advertisement"""
        current_ad = self.ads[self.current_ad_index]
        webbrowser.open_new(current_ad["url"])
        
    def create_main_menu(self):
        menubar = Menu(self.root)
        self.root.config(menu=menubar)

        # File Menu
        file_menu = Menu(menubar, tearoff=0)
        menubar.add_cascade(label="File", menu=file_menu)
        file_menu.add_command(label="Save Project", command=self.save_project)
        file_menu.add_command(label="Load Project", command=self.load_project)
        file_menu.add_separator()
        file_menu.add_command(label="Export to Excel", command=self.export_to_excel)
        file_menu.add_separator()
        file_menu.add_command(label="Exit", command=self.root.quit)

        # View Menu for Themes
        view_menu = Menu(menubar, tearoff=0)
        menubar.add_cascade(label="View", menu=view_menu)
        theme_menu = Menu(view_menu, tearoff=0)
        view_menu.add_cascade(label="Themes", menu=theme_menu)

        for theme in sorted(self.root.get_themes()):
            theme_menu.add_command(label=theme, command=lambda t=theme: self.change_theme(t))
            
    def change_theme(self, theme_name):
        try:
            self.root.set_theme(theme_name)
            self.app_theme = theme_name
            self.save_settings()
            self.show_floating_message(f"Theme changed to {theme_name}")
        except Exception as e:
            self.show_floating_message(f"Failed to set theme: {e}")

    def _open_website_link(self, event):
        webbrowser.open_new("http://intellsoft.ir")

    def _on_link_enter(self, event):
        self.footer_text_widget.config(cursor="hand2")

    def _on_link_leave(self, event):
        self.footer_text_widget.config(cursor="arrow")

    def start_bonjour_scan(self):
        """Starts Bonjour device discovery"""
        if not self.zeroconf:
            try:
                self.zeroconf = Zeroconf()
                self.bonjour_listener = BonjourListener(self)
                self.browser = ServiceBrowser(
                    self.zeroconf, 
                    ["_http._tcp.local.", "_onvif._tcp.local."],
                    self.bonjour_listener
                )
                self.show_floating_message("Bonjour scan started")
            except Exception as e:
                self.show_floating_message(f"Bonjour error: {str(e)}")
            
    def stop_bonjour_scan(self):
        """Stops Bonjour device discovery"""
        if self.zeroconf:
            self.zeroconf.close()
            self.zeroconf = None
            self.show_floating_message("Bonjour scan stopped")
            
    def add_bonjour_device(self, ip, port):
        """Adds a Bonjour device to the table"""
        existing = False
        for child in self.tree.get_children():
            if self.tree.item(child, 'values')[0] == ip:
                existing = True
                break
        
        if not existing:
            device = {
                "IP": ip,
                "MAC": "N/A (Bonjour)",
                "Credential Set": "Default Admin",
                "Open Ports": str(port),
                "RTSP URL": "",
                "Camera Status": "Bonjour Device",
                "Snapshot": "Unavailable"
            }
            item_id = self.add_device_to_tree(device)
            self.root.after(0, self.update_status, f"Bonjour device found: ({ip})")
            self.onvif_queue.put((ip, item_id))
            self.port_scan_queue.put((ip, item_id))

    def on_right_click(self, event):
        item_id = self.tree.identify_row(event.y)
        if not item_id: return
        
        column_id = self.tree.identify_column(event.x)
        if not column_id: return

        col_index = int(column_id.replace('#', '')) - 1
        columns = ('IP', 'MAC', 'Credential Set', 'Open Ports', 'RTSP URL', 'Camera Status', 'Snapshot')
        column_name = columns[col_index] if 0 <= col_index < len(columns) else None

        if not column_name: return

        cell_value = self.tree.item(item_id, 'values')[col_index]

        context_menu = Menu(self.root, tearoff=0)
        context_menu.add_command(label=f"Copy '{cell_value}'", command=lambda: self.copy_to_clipboard(cell_value))
        context_menu.post(event.x_root, event.y_root)

    def copy_to_clipboard(self, text):
        self.root.clipboard_clear()
        self.root.clipboard_append(text)
        self.show_floating_message(f"'{text}' copied to clipboard")

    def export_to_excel(self):
        if not self.tree.get_children():
            self.show_floating_message("No data to export")
            return

        file_path = filedialog.asksaveasfilename(
            defaultextension=".xlsx",
            filetypes=[("Excel files", "*.xlsx"), ("All files", "*.*")],
            title="Save as Excel File"
        )
        if not file_path:
            return

        try:
            workbook = openpyxl.Workbook()
            sheet = workbook.active
            sheet.title = "Network Scan Results"

            headers = [self.tree.heading(col)['text'] for col in self.tree['columns']]
            sheet.append(headers)

            for item_id in self.tree.get_children(): 
                values = self.tree.item(item_id, 'values')
                sheet.append(values)

            workbook.save(file_path)
            self.show_floating_message(f"Data exported to {file_path}")
        except Exception as e:
            self.show_floating_message(f"Error saving Excel: {e}")

    def update_cred_combobox_options(self):
        self.cred_set_combobox['values'] = list(self.credential_sets.keys())
        if self.cred_set_combobox['values']:
            self.cred_set_var.set(self.cred_set_combobox['values'][0])

    def on_cred_set_selected(self, event):
        pass

    def apply_camera_credentials(self):
        selected = self.tree.selection()
        if not selected:
            self.show_floating_message("Please select a camera")
            return
            
        item = selected[0]
        values = self.tree.item(item, 'values')
        ip = values[0]
        
        selected_cred_set_name = self.cred_set_var.get()
        if not selected_cred_set_name or selected_cred_set_name not in self.credential_sets:
            self.show_floating_message("Please select a valid credential set")
            return

        self.camera_ip_to_cred_set[ip] = selected_cred_set_name
        self.save_camera_data()
            
        new_values = list(values)
        if len(new_values) >= 3: 
            new_values[2] = selected_cred_set_name
            self.tree.item(item, values=new_values)
            self.onvif_queue.put((ip, item))
            self.show_floating_message("Credential set applied")
        else:
            self.show_floating_message("Invalid data structure")

    def on_tree_select(self, event):
        selected = self.tree.selection()
        if selected:
            item = selected[0]
            values = self.tree.item(item, 'values')
            ip = values[0]
            associated_cred_set = self.camera_ip_to_cred_set.get(ip, "Default Admin")
            self.cred_set_var.set(associated_cred_set)

    def open_credential_manager(self):
        cred_manager_window = tk.Toplevel(self.root)
        cred_manager_window.title("Manage Credential Sets")
        self.center_window(cred_manager_window)
        cred_manager_window.transient(self.root)
        cred_manager_window.grab_set()

        input_frame = ttk.Frame(cred_manager_window, padding="10")
        input_frame.pack(fill=tk.X)

        ttk.Label(input_frame, text="Set Name:").grid(row=0, column=0, padx=5, pady=5, sticky="w")
        cred_name_var = tk.StringVar()
        cred_name_entry = ttk.Entry(input_frame, textvariable=cred_name_var, width=30)
        cred_name_entry.grid(row=0, column=1, padx=5, pady=5, sticky="ew")

        ttk.Label(input_frame, text="Username:").grid(row=1, column=0, padx=5, pady=5, sticky="w")
        cred_username_var = tk.StringVar()
        cred_username_entry = ttk.Entry(input_frame, textvariable=cred_username_var, width=30)
        cred_username_entry.grid(row=1, column=1, padx=5, pady=5, sticky="ew")

        ttk.Label(input_frame, text="Password:").grid(row=2, column=0, padx=5, pady=5, sticky="w")
        cred_password_var = tk.StringVar()
        cred_password_entry = ttk.Entry(input_frame, textvariable=cred_password_var, width=30, show='*')
        cred_password_entry.grid(row=2, column=1, padx=5, pady=5, sticky="ew")

        button_frame = ttk.Frame(cred_manager_window, padding="10")
        button_frame.pack(fill=tk.X)

        def add_cred_set():
            name = cred_name_var.get().strip()
            username = cred_username_var.get().strip()
            password = cred_password_var.get()

            if not name or not username:
                self.show_floating_message("Set name and username cannot be empty")
                return
            if name in self.credential_sets and name != "Default Admin":
                self.show_floating_message(f"Set '{name}' already exists")
                return
            
            encrypted_password = self.encrypt_password(password)
            self.credential_sets[name] = {"username": username, "password": encrypted_password}
            self.save_camera_data()
            update_cred_treeview()
            self.update_cred_combobox_options()
            self.show_floating_message(f"Credential set '{name}' added")
            cred_name_var.set("")
            cred_username_var.set("")
            cred_password_var.set("")

        def update_cred_set():
            selected = cred_tree.selection()
            if not selected:
                self.show_floating_message("Please select a credential set")
                return
            
            old_name = cred_tree.item(selected[0], 'values')[0]
            new_name = cred_name_var.get().strip()
            username = cred_username_var.get().strip()
            password = cred_password_var.get()

            if not new_name or not username:
                self.show_floating_message("Set name and username cannot be empty")
                return
            
            if old_name == "Default Admin" and new_name != "Default Admin":
                self.show_floating_message("Cannot change 'Default Admin' name")
                return

            if new_name != old_name and new_name in self.credential_sets:
                self.show_floating_message(f"Set '{new_name}' already exists")
                return

            encrypted_password = self.encrypt_password(password)

            if new_name != old_name:
                del self.credential_sets[old_name]
                for ip, name in list(self.camera_ip_to_cred_set.items()):
                    if name == old_name:
                        self.camera_ip_to_cred_set[ip] = new_name

            self.credential_sets[new_name] = {"username": username, "password": encrypted_password}
            self.save_camera_data()
            update_cred_treeview()
            self.update_cred_combobox_options()
            self.show_floating_message(f"Credentials '{old_name}' updated")
            cred_name_var.set("")
            cred_username_var.set("")
            cred_password_var.set("")

        def delete_cred_set():
            selected = cred_tree.selection()
            if not selected:
                self.show_floating_message("Please select a credential set")
                return
            
            name_to_delete = cred_tree.item(selected[0], 'values')[0]
            
            if name_to_delete == "Default Admin":
                self.show_floating_message("Cannot delete 'Default Admin'")
                return

            cameras_using_this_set = [ip for ip, name in self.camera_ip_to_cred_set.items() if name == name_to_delete]
            if cameras_using_this_set:
                confirm = messagebox.askyesno(
                    "Confirm Delete", 
                    f"Credential set '{name_to_delete}' is used by {len(cameras_using_this_set)} cameras.\nAre you sure you want to delete it? Affected cameras will revert to 'Default Admin'.", 
                    parent=cred_manager_window
                )
                if not confirm:
                    return
                for ip in cameras_using_this_set:
                    self.camera_ip_to_cred_set[ip] = "Default Admin"
                    for item_id in self.tree.get_children(): 
                        if self.tree.item(item_id, 'values')[0] == ip:
                            current_values = list(self.tree.item(item_id, 'values'))
                            current_values[2] = "Default Admin"
                            self.tree.item(item_id, values=current_values)

            del self.credential_sets[name_to_delete]
            self.save_camera_data()
            update_cred_treeview()
            self.update_cred_combobox_options()
            self.show_floating_message(f"Credentials '{name_to_delete}' deleted")
            cred_name_var.set("")
            cred_username_var.set("")
            cred_password_var.set("")

        ttk.Button(button_frame, text="Add", command=add_cred_set).pack(side=tk.LEFT, padx=5, pady=5)
        ttk.Button(button_frame, text="Update", command=update_cred_set).pack(side=tk.LEFT, padx=5, pady=5)
        ttk.Button(button_frame, text="Delete", command=delete_cred_set).pack(side=tk.LEFT, padx=5, pady=5)

        cred_tree_frame = ttk.Frame(cred_manager_window, padding="10")
        cred_tree_frame.pack(fill=tk.BOTH, expand=True)

        cred_columns = ('Set Name', 'Username', 'Password')
        cred_tree = ttk.Treeview(cred_tree_frame, columns=cred_columns, show='headings')
        for col in cred_columns:
            cred_tree.heading(col, text=col)
            cred_tree.column(col, width=150)

        cred_scrollbar = ttk.Scrollbar(cred_tree_frame, orient=tk.VERTICAL, command=cred_tree.yview)
        cred_tree.configure(yscroll=cred_scrollbar.set)
        cred_scrollbar.pack(side=tk.RIGHT, fill=tk.Y)
        cred_tree.pack(fill=tk.BOTH, expand=True)

        def on_cred_tree_select(event):
            selected = cred_tree.selection()
            if selected:
                values = cred_tree.item(selected[0], 'values')
                cred_name_var.set(values[0])
                cred_username_var.set(values[1])
                cred_password_var.set("") 

        cred_tree.bind('<<TreeviewSelect>>', on_cred_tree_select)

        def update_cred_treeview():
            for item in cred_tree.get_children():
                cred_tree.delete(item)
            for name, creds in self.credential_sets.items():
                decrypted_password_display = self.decrypt_password(creds['password'])
                cred_tree.insert('', tk.END, values=(name, creds['username'], decrypted_password_display))

        update_cred_treeview()

        close_button_frame = ttk.Frame(cred_manager_window, padding="10")
        close_button_frame.pack(fill=tk.X)
        ttk.Button(close_button_frame, text="Close", command=cred_manager_window.destroy).pack(side=tk.RIGHT, padx=5, pady=5)

        cred_manager_window.protocol("WM_DELETE_WINDOW", cred_manager_window.destroy)
        cred_manager_window.wait_window()

    def open_scan_range_manager(self):
        range_manager_window = tk.Toplevel(self.root)
        range_manager_window.title("Manage Scan Range")
        self.center_window(range_manager_window)
        range_manager_window.transient(self.root)
        range_manager_window.grab_set()

        input_frame = ttk.Frame(range_manager_window, padding="10")
        input_frame.pack(fill=tk.X)

        ttk.Label(input_frame, text="Start IP Address:").grid(row=0, column=0, padx=5, pady=5, sticky="w")
        self.start_ip_var = tk.StringVar(value=self.scan_range.get("start_ip", ""))
        start_ip_entry = ttk.Entry(input_frame, textvariable=self.start_ip_var, width=20)
        start_ip_entry.grid(row=0, column=1, padx=5, pady=5, sticky="ew")

        ttk.Label(input_frame, text="End IP Address:").grid(row=1, column=0, padx=5, pady=5, sticky="w")
        self.end_ip_var = tk.StringVar(value=self.scan_range.get("end_ip", ""))
        end_ip_entry = ttk.Entry(input_frame, textvariable=self.end_ip_var, width=20)
        end_ip_entry.grid(row=1, column=1, padx=5, pady=5, sticky="ew")

        button_frame = ttk.Frame(range_manager_window, padding="10")
        button_frame.pack(fill=tk.X)

        def save_range():
            start_ip = self.start_ip_var.get().strip()
            end_ip = self.end_ip_var.get().strip()

            if not start_ip and not end_ip:
                self.scan_range = {"start_ip": "", "end_ip": ""}
                self.save_settings()
                self.show_floating_message("Scan range cleared")
                return

            try:
                start_ip_obj = ipaddress.IPv4Address(start_ip)
                end_ip_obj = ipaddress.IPv4Address(end_ip)
                if start_ip_obj > end_ip_obj:
                    self.show_floating_message("Invalid IP range")
                    return
                self.scan_range = {"start_ip": start_ip, "end_ip": end_ip}
                self.save_settings()
                self.show_floating_message("Scan range saved")
            except ipaddress.AddressValueError:
                self.show_floating_message("Invalid IP format")
            except Exception as e:
                self.show_floating_message(f"Error: {str(e)}")

        ttk.Button(button_frame, text="Save Range", command=save_range).pack(side=tk.LEFT, padx=5, pady=5)
        ttk.Button(button_frame, text="Clear Range", command=lambda: [self.start_ip_var.set(""), self.end_ip_var.set(""), save_range()]).pack(side=tk.LEFT, padx=5, pady=5)
        ttk.Button(button_frame, text="Close", command=range_manager_window.destroy).pack(side=tk.RIGHT, padx=5, pady=5)

        range_manager_window.protocol("WM_DELETE_WINDOW", range_manager_window.destroy)
        range_manager_window.wait_window()

    # FIXED: Added CMD Tools dropdown menu functionality
    def _create_cmd_tool_menus(self):
        """Creates the dropdown menus for CMD Tools."""
        # Network and Troubleshooting (Priority 1)
        network_menu = Menu(self.cmd_tools_menu, tearoff=0)
        self.cmd_tools_menu.add_cascade(label="Network & Troubleshooting", menu=network_menu)
        
        # Sub-menu: IP Configuration
        ip_config_menu = Menu(network_menu, tearoff=0)
        network_menu.add_cascade(label="IP Configuration", menu=ip_config_menu)
        ip_config_menu.add_command(label="ipconfig", command=lambda: self.open_cmd_tools_window("ipconfig"))
        ip_config_menu.add_command(label="ipconfig /all", command=lambda: self.open_cmd_tools_window("ipconfig /all"))
        ip_config_menu.add_command(label="ipconfig /renew", command=lambda: self.open_cmd_tools_window("ipconfig /renew"))
        ip_config_menu.add_command(label="ipconfig /release", command=lambda: self.open_cmd_tools_window("ipconfig /release"))

        # Sub-menu: Connectivity Test
        conn_test_menu = Menu(network_menu, tearoff=0)
        network_menu.add_cascade(label="Connectivity Test", menu=conn_test_menu)
        conn_test_menu.add_command(label="ping google.com", command=lambda: self.open_cmd_tools_window("ping google.com -n 4"))
        conn_test_menu.add_command(label="ping [IP/Hostname] -t", command=lambda: self.open_cmd_tools_window("ping [Enter IP/Hostname] -t", show_input=True))
        conn_test_menu.add_command(label="tracert google.com", command=lambda: self.open_cmd_tools_window("tracert google.com"))
        conn_test_menu.add_command(label="pathping google.com", command=lambda: self.open_cmd_tools_window("pathping google.com"))
        conn_test_menu.add_command(label="telnet [IP] [Port]", command=lambda: self.open_cmd_tools_window("telnet [Enter IP] [Enter Port]", show_input=True))


        # Sub-menu: Network Connections & Ports
        net_conn_menu = Menu(network_menu, tearoff=0)
        network_menu.add_cascade(label="Network Connections & Ports", menu=net_conn_menu)
        net_conn_menu.add_command(label="netstat -an", command=lambda: self.open_cmd_tools_window("netstat -an"))
        net_conn_menu.add_command(label="nslookup google.com", command=lambda: self.open_cmd_tools_window("nslookup google.com"))
        net_conn_menu.add_command(label="getmac", command=lambda: self.open_cmd_tools_window("getmac"))
        net_conn_menu.add_command(label="arp -a", command=lambda: self.open_cmd_tools_window("arp -a"))

        # Sub-menu: Network Configuration
        net_config_menu = Menu(network_menu, tearoff=0)
        network_menu.add_cascade(label="Network Configuration", menu=net_config_menu)
        net_config_menu.add_command(label="netsh (Interactive)", command=lambda: self.open_cmd_tools_window("netsh")) # User will need to interact
        net_config_menu.add_command(label="route print", command=lambda: self.open_cmd_tools_window("route print"))

        # System Information & Process Management (Priority 2)
        system_menu = Menu(self.cmd_tools_menu, tearoff=0)
        self.cmd_tools_menu.add_cascade(label="System & Processes", menu=system_menu)
        system_menu.add_command(label="systeminfo", command=lambda: self.open_cmd_tools_window("systeminfo"))
        system_menu.add_command(label="tasklist", command=lambda: self.open_cmd_tools_window("tasklist"))
        system_menu.add_command(label="taskkill /F /IM [Image Name]", command=lambda: self.open_cmd_tools_window("taskkill /F /IM [Enter Image Name]", show_input=True))
        system_menu.add_command(label="sc query [Service Name]", command=lambda: self.open_cmd_tools_window("sc query [Enter Service Name]", show_input=True))
        system_menu.add_command(label="driverquery", command=lambda: self.open_cmd_tools_window("driverquery"))
        system_menu.add_command(label="powercfg /energy", command=lambda: self.open_cmd_tools_window("powercfg /energy"))

        # Add separator and full window option
        self.cmd_tools_menu.add_separator()
        self.cmd_tools_menu.add_command(label="Open Full CMD Tools Window", command=self.open_cmd_tools_window)

    # FIXED: Added show_input parameter to handle commands requiring input
    def open_cmd_tools_window(self, predefined_command="", show_input=False):
        """
        Opens a window containing CMD network commands.
        Now accepts a predefined command and shows input field if needed.
        """
        cmd_window = tk.Toplevel(self.root)
        cmd_window.title("CMD Tools")
        self.center_window(cmd_window)
        cmd_window.transient(self.root)
        cmd_window.grab_set()
        
        # Set minimum size
        cmd_window.minsize(600, 400)
        
        # Main frame
        main_frame = ttk.Frame(cmd_window, padding="10")
        main_frame.pack(fill=tk.BOTH, expand=True)
        
        # Command entry frame
        cmd_entry_frame = ttk.Frame(main_frame)
        cmd_entry_frame.pack(fill=tk.X, pady=5)
        
        ttk.Label(cmd_entry_frame, text="Command:").pack(side=tk.LEFT, padx=5)
        
        cmd_entry_var = tk.StringVar(value=predefined_command)
        cmd_entry = ttk.Entry(cmd_entry_frame, textvariable=cmd_entry_var, width=70)
        cmd_entry.pack(side=tk.LEFT, fill=tk.X, expand=True, padx=5)
        
        # If show_input is True, replace placeholders with actual input fields
        if show_input:
            cmd_entry.bind("<FocusIn>", lambda e: self._handle_command_placeholders(cmd_entry, cmd_entry_var))
        
        execute_btn = ttk.Button(
            cmd_entry_frame, 
            text="Execute", 
            command=lambda: self._execute_cmd_tool(cmd_output_text, cmd_entry_var.get())
        )
        execute_btn.pack(side=tk.LEFT, padx=5)
        
        # Output frame
        output_frame = ttk.LabelFrame(main_frame, text="Output")
        output_frame.pack(fill=tk.BOTH, expand=True, pady=5)
        
        # Create text widget with scrollbar
        cmd_output_text = tk.Text(output_frame, wrap=tk.WORD, state=tk.DISABLED)
        cmd_output_text.pack(fill=tk.BOTH, expand=True, padx=5, pady=5)
        
        scrollbar = ttk.Scrollbar(output_frame, command=cmd_output_text.yview)
        scrollbar.pack(side=tk.RIGHT, fill=tk.Y)
        cmd_output_text.config(yscrollcommand=scrollbar.set)
        
        # Button frame
        button_frame = ttk.Frame(main_frame)
        button_frame.pack(fill=tk.X, pady=5)
        
        clear_btn = ttk.Button(
            button_frame,
            text="Clear Output",
            command=lambda: self._clear_cmd_output(cmd_output_text)
        )
        clear_btn.pack(side=tk.LEFT, padx=5)
        
        copy_btn = ttk.Button(
            button_frame,
            text="Copy Output",
            command=lambda: self._copy_cmd_output(cmd_output_text)
        )
        copy_btn.pack(side=tk.LEFT, padx=5)
        
        close_btn = ttk.Button(
            button_frame,
            text="Close",
            command=cmd_window.destroy
        )
        close_btn.pack(side=tk.RIGHT, padx=5)
        
        # Focus on entry field
        cmd_entry.focus_set()
        
        cmd_window.protocol("WM_DELETE_WINDOW", cmd_window.destroy)

    def _handle_command_placeholders(self, cmd_entry, cmd_entry_var):
        """Replaces placeholders in command with actual input fields."""
        command = cmd_entry_var.get()
        if "[Enter" in command:
            # Create a dialog to get the values
            input_dialog = tk.Toplevel(self.root)
            input_dialog.title("Enter Command Parameters")
            self.center_window(input_dialog)
            input_dialog.transient(self.root)
            input_dialog.grab_set()
            
            # Find all placeholders
            placeholders = re.findall(r'\[([^\]]+)\]', command)
            entries = {}
            row = 0
            
            for placeholder in placeholders:
                ttk.Label(input_dialog, text=f"{placeholder}:").grid(row=row, column=0, padx=5, pady=2, sticky="e")
                entry_var = tk.StringVar()
                entry = ttk.Entry(input_dialog, textvariable=entry_var, width=30)
                entry.grid(row=row, column=1, padx=5, pady=2, sticky="w")
                entries[placeholder] = entry_var
                row += 1
            
            def apply_values():
                new_command = command
                for placeholder, var in entries.items():
                    value = var.get()
                    if value:
                        new_command = new_command.replace(f"[{placeholder}]", value)
                cmd_entry_var.set(new_command)
                input_dialog.destroy()
                cmd_entry.focus_set()
            
            button_frame = ttk.Frame(input_dialog)
            button_frame.grid(row=row, column=0, columnspan=2, pady=10)
            
            ttk.Button(button_frame, text="Apply", command=apply_values).pack(side=tk.LEFT, padx=5)
            ttk.Button(button_frame, text="Cancel", command=input_dialog.destroy).pack(side=tk.LEFT, padx=5)
            
            input_dialog.wait_window()

    def _execute_cmd_tool(self, output_widget, command):
        """Executes a command and displays output."""
        if not command:
            self.show_floating_message("Please enter a command")
            return
            
        output_widget.config(state=tk.NORMAL)
        output_widget.delete(1.0, tk.END)
        output_widget.insert(tk.END, f"Executing: {command}\n\n")
        output_widget.see(tk.END)
        output_widget.config(state=tk.DISABLED)
        
        # Execute command in a separate thread
        threading.Thread(
            target=self._run_cmd_in_thread, 
            args=(output_widget, command),
            daemon=True
        ).start()

    def _run_cmd_in_thread(self, output_widget, command):
        """Runs command in a background thread and updates UI."""
        try:
            creation_flags = subprocess.CREATE_NO_WINDOW if platform.system() == "Windows" else 0
            process = subprocess.run(
                command,
                shell=True,
                capture_output=True,
                text=True,
                check=False,
                encoding='utf-8',
                errors='replace',
                creationflags=creation_flags,
                timeout=60
            )
            output = process.stdout
            error_output = process.stderr
            
            # Update UI in main thread
            self.root.after(0, self._update_cmd_output, output_widget, output, error_output)
        except subprocess.TimeoutExpired:
            self.root.after(0, self._update_cmd_output, output_widget, "", f"Command timed out after 60 seconds", True)
        except Exception as e:
            self.root.after(0, self._update_cmd_output, output_widget, "", f"Error: {str(e)}", True)

    def _update_cmd_output(self, output_widget, output, error_output="", is_error=False):
        """Updates the command output widget."""
        output_widget.config(state=tk.NORMAL)
        output_widget.insert(tk.END, output)
        
        if error_output:
            output_widget.insert(tk.END, f"\n\n--- ERROR ---\n{error_output}")
        
        if is_error:
            output_widget.tag_configure("error", foreground="red")
            output_widget.insert(tk.END, "\n\n--- COMMAND FAILED ---\n", "error")
        
        output_widget.see(tk.END)
        output_widget.config(state=tk.DISABLED)

    def _clear_cmd_output(self, output_widget):
        """Clears the command output."""
        output_widget.config(state=tk.NORMAL)
        output_widget.delete(1.0, tk.END)
        output_widget.config(state=tk.DISABLED)

    def _copy_cmd_output(self, output_widget):
        """Copies command output to clipboard."""
        output_widget.config(state=tk.NORMAL)
        content = output_widget.get(1.0, tk.END)
        self.root.clipboard_clear()
        self.root.clipboard_append(content)
        output_widget.config(state=tk.DISABLED)
        self.show_floating_message("Output copied to clipboard")

    def start_search(self):
        if self.scanning_active:
            self.show_floating_message("Scan already active.")
            return

        self.stop_flag = False
        self.scanning_active = True
        self.found_count = 0
        self.tree.delete(*self.tree.get_children())
        self.progress_bar.pack(side=tk.RIGHT, fill=tk.X, expand=True)
        self.update_status("Starting network scan...")
        self.search_btn.config(state=tk.DISABLED)
        self.stop_btn.config(state=tk.NORMAL)
        
        self.search_thread = threading.Thread(target=self._run_search_in_thread, daemon=True)
        self.search_thread.start()

    def stop_search(self):
        self.stop_flag = True
        self.scanning_active = False
        self.update_status("Scan stopped.")
        self.search_btn.config(state=tk.NORMAL)
        self.stop_btn.config(state=tk.DISABLED)
        self.progress_bar['value'] = 0
        self.progress_bar.pack_forget()

    def _run_search_in_thread(self):
        start_ip_str = self.scan_range.get("start_ip")
        end_ip_str = self.scan_range.get("end_ip")

        if not start_ip_str or not end_ip_str:
            self.root.after(0, self.show_floating_message, "Scan range not set. Please go to 'Manage Scan Range'.")
            self.root.after(0, self.stop_search)
            return

        try:
            start_ip = int(ipaddress.IPv4Address(start_ip_str))
            end_ip = int(ipaddress.IPv4Address(end_ip_str))
        except ipaddress.AddressValueError:
            self.root.after(0, self.show_floating_message, "Invalid IP address format in scan range.")
            self.root.after(0, self.stop_search)
            return

        total_ips = end_ip - start_ip + 1
        scanned_ips = 0

        max_workers = 100 
        with concurrent.futures.ThreadPoolExecutor(max_workers=max_workers) as executor:
            futures = {executor.submit(self._ping_ip, str(ipaddress.IPv4Address(i))): str(ipaddress.IPv4Address(i)) for i in range(start_ip, end_ip + 1)}
            
            for future in concurrent.futures.as_completed(futures):
                if self.stop_flag:
                    break
                
                scanned_ip_str = futures[future]
                scanned_ips += 1
                progress_value = (scanned_ips / total_ips) * 100
                self.root.after(0, self.update_progress, progress_value, f"Scanning: {scanned_ip_str}")
                
                try:
                    mac_address = future.result()
                    if mac_address:
                        device_data = {
                            "IP": scanned_ip_str,
                            "MAC": mac_address,
                            "Credential Set": "Default Admin",
                            "Open Ports": "Scanning...",
                            "RTSP URL": "N/A",
                            "Camera Status": "Searching...",
                            "Snapshot": "N/A"
                        }
                        self.root.after(0, self.add_device_to_tree, device_data)
                except Exception as e:
                    print(f"Error processing IP {scanned_ip_str}: {e}")

        self.root.after(0, self.finalize_scan)

    def _ping_ip(self, ip_address):
        """Pings an IP address and returns MAC if successful, None otherwise."""
        try:
            response = ping(ip_address, timeout=0.5, unit='ms')
            if response is not False and response is not None:
                mac = get_mac_address(ip=ip_address)
                return mac if mac else "Unknown"
        except Exception:
            pass
        return None

    def get_hostname(self, ip_address):
        """Attempts to resolve hostname from IP address."""
        try:
            return socket.gethostbyaddr(ip_address)[0]
        except socket.herror:
            return "N/A"

    def get_mac_vendor(self, mac_address):
        return "N/A"

    def add_device_to_tree(self, device_data):
        """
        Adds a device to the treeview.
        """
        self.found_count += 1
        ip = device_data["IP"]
        mac = device_data["MAC"]
        cred_set = device_data["Credential Set"]
        open_ports = device_data["Open Ports"]
        rtsp_url = device_data["RTSP URL"]
        camera_status = device_data["Camera Status"]
        snapshot = device_data["Snapshot"]

        item_id = self.tree.insert(
            '', tk.END, 
            values=(ip, mac, cred_set, open_ports, rtsp_url, camera_status, snapshot),
            tags=('ip_dup',) if self.is_ip_duplicate(ip) else ()
        )

        self.update_status(f"Found {self.found_count} devices.")
        self.onvif_queue.put((ip, item_id))
        self.port_scan_queue.put((ip, item_id))
    
    def update_camera_info(self, item_id, cred_set_name, rtsp_url, camera_status, snapshot_uri):
        """Updates ONVIF camera specific info in the tree."""
        if not self.tree.exists(item_id): return

        current_values = list(self.tree.item(item_id, 'values'))
        
        if len(current_values) >= 7:
            current_values[2] = cred_set_name
            current_values[4] = rtsp_url
            current_values[5] = camera_status
            
            # Store snapshot_uri in tags and set appropriate text
            current_tags = list(self.tree.item(item_id, 'tags'))
            # Remove any existing snapshot_uri tag
            current_tags = [tag for tag in current_tags if not tag.startswith('snapshot_uri:')]
            if snapshot_uri:
                current_tags.append(f'snapshot_uri:{snapshot_uri}')
            
            # Set snapshot column text based on availability
            snapshot_status_text = "View Image" if snapshot_uri or rtsp_url else "Unavailable"
            current_values[6] = snapshot_status_text
            
            # Update camera status tags
            if camera_status == "ONVIF Found" and 'camera_found' not in current_tags:
                current_tags.append('camera_found')
            elif camera_status != "ONVIF Found" and 'camera_found' in current_tags:
                if 'camera_found' in current_tags:
                    current_tags.remove('camera_found')
            
            if camera_status in ["Auth Failed", "Timeout", "Error"] and 'camera_error' not in current_tags:
                current_tags.append('camera_error')
            elif camera_status not in ["Auth Failed", "Timeout", "Error"] and 'camera_error' in current_tags:
                if 'camera_error' in current_tags:
                    current_tags.remove('camera_error')
            
            # Update the tree item
            self.tree.item(item_id, values=current_values, tags=current_tags)
        else:
            print("Invalid values length in update_camera_info")

    def update_open_ports(self, item_id, open_ports):
        """Updates the Open Ports column in the treeview."""
        if not self.tree.exists(item_id): return
        current_values = list(self.tree.item(item_id, 'values'))
        if len(current_values) >= 4:
            current_values[3] = ", ".join(map(str, open_ports)) if open_ports else "N/A"
        self.tree.item(item_id, values=current_values)

    def is_ip_duplicate(self, ip):
        """Checks if an IP already exists in the treeview."""
        for item_id in self.tree.get_children():
            if self.tree.item(item_id, 'values')[0] == ip:
                return True
        return False

    def get_child_items(self, parent_item):
        """Helper to get all children of a parent, including nested if any."""
        children = self.tree.get_children(parent_item)
        all_children = list(children)
        for child in children:
            all_children.extend(self.get_child_items(child))
        return all_children

    def sort_treeview_column(self, col, reverse):
        data = [(self.tree.set(child, col), child) for child in self.tree.get_children('')]
        
        # Basic sorting for now, can be improved for numeric/IP sorting
        data.sort(reverse=reverse)
        
        for index, (val, item) in enumerate(data):
            self.tree.move(item, '', index)
        self.tree.heading(col, command=lambda: self.sort_treeview_column(col, not reverse))

    def update_status(self, message):
        self.progress_label.config(text=message)

    def update_progress(self, value, message):
        self.progress_bar['value'] = value
        self.update_status(message)

    def finalize_scan(self):
        self.stop_search()
        self.update_status(f"Scan finished. Found {self.found_count} devices.")
        self.progress_bar['value'] = 100
        self.show_floating_message(f"Scan finished. Found {self.found_count} devices.")

    def filter_results(self, event=None):
        query = self.search_var.get().lower()
        
        # First, re-attach all items to make sure we search everything
        all_items = self.tree.get_children('')
        for item_id in all_items:
             self.tree.reattach(item_id, '', tk.END)

        if not query:
            return

        # Now, detach items that do not match
        for item_id in self.tree.get_children(''):
            values = [str(v).lower() for v in self.tree.item(item_id, 'values')]
            
            if not any(query in v for v in values):
                self.tree.detach(item_id)

    def _on_tree_motion(self, event):
        """Handles mouse motion over the Treeview for hover functionality."""
        item_id = self.tree.identify_row(event.y)
        column_id = self.tree.identify_column(event.x)

        # Check if we are over the Snapshot column (column #7)
        if item_id and column_id == '#7':
            item_values = self.tree.item(item_id, 'values')
            snapshot_status_text = item_values[6]  # 6th index in the values tuple

            if snapshot_status_text == "View Image":
                if item_id != self.current_hover_item:
                    self._hide_hover_image()
                    self.current_hover_item = item_id

                    # Extract the snapshot_uri from the tags
                    snapshot_uri = ""
                    item_tags = self.tree.item(item_id, 'tags')
                    for tag in item_tags:
                        if tag.startswith('snapshot_uri:'):
                            snapshot_uri = tag.split(':', 1)[1]
                            break

                    # Get the RTSP URL and credentials
                    ip = item_values[0]
                    cred_set_name = self.camera_ip_to_cred_set.get(ip, "Default Admin")
                    cred_info = self.credential_sets.get(cred_set_name)
                    username = cred_info['username'] if cred_info else "admin"
                    password = self.decrypt_password(cred_info['password']) if cred_info else "admin"
                    rtsp_url = item_values[4]  # RTSP URL is at index 4

                    # Create hover window
                    self.hover_window = tk.Toplevel(self.root)
                    self.hover_window.wm_overrideredirect(True)
                    self.hover_window.wm_geometry(f"+{event.x_root + 10}+{event.y_root + 10}")
                    self.hover_window.attributes("-topmost", True)

                    self.hover_image_label = ttk.Label(self.hover_window, text="Loading...")
                    self.hover_image_label.pack(padx=5, pady=5)

                    self.last_hover_thread_id += 1
                    thread_id = self.last_hover_thread_id
                    self.hover_image_thread = threading.Thread(
                        target=self._fetch_thumbnail_for_hover,
                        args=(thread_id, ip, username, password, rtsp_url, snapshot_uri)
                    )
                    self.hover_image_thread.daemon = True
                    self.hover_image_thread.start()
            else:
                self._hide_hover_image()
        else:
            self._hide_hover_image()

    def _on_tree_leave(self, event):
        """Handles mouse leaving the Treeview, hides hover image."""
        self._hide_hover_image()

    def _hide_hover_image(self):
        """Hides and destroys the hover image window."""
        if self.hover_window:
            self.hover_window.destroy()
            self.hover_window = None
            self.hover_image_label = None
            self.current_hover_item = None
            self.current_hover_image_tk = None

    def _fetch_thumbnail_for_hover(self, thread_id, ip, username, password, rtsp_url, snapshot_uri):
        """Fetches and processes a thumbnail for hover display."""
        if thread_id != self.last_hover_thread_id:
            return

        thumbnail_size = (160, 120)
        image_key = f"hover_thumbnail_{ip}"

        # Check cache first
        if image_key in self.hover_image_cache:
            self.root.after(0, lambda: self._update_hover_window(thread_id, self.hover_image_cache[image_key]))
            return

        image_data = None
        
        # First try snapshot URI if available
        if snapshot_uri:
            try:
                auth = (username, password) if username and password else None
                response = requests.get(snapshot_uri, auth=auth, timeout=3)
                response.raise_for_status()
                image_data = response.content
            except Exception as e:
                print(f"Error fetching thumbnail from Snapshot URI ({snapshot_uri}): {e}")
                image_data = None
        
        # If snapshot URI fails, try RTSP stream
        if not image_data and rtsp_url:
            try:
                # Add credentials to RTSP URL if needed
                if username and password:
                    parsed_url = urlparse(rtsp_url)
                    netloc_with_auth = f"{username}:{password}@{parsed_url.hostname}"
                    if parsed_url.port:
                        netloc_with_auth += f":{parsed_url.port}"
                    rtsp_url_with_auth = parsed_url._replace(netloc=netloc_with_auth).geturl()
                else:
                    rtsp_url_with_auth = rtsp_url
                
                # Capture frame from RTSP stream
                cap = cv2.VideoCapture(rtsp_url_with_auth)
                if not cap.isOpened():
                    raise ConnectionError("Cannot open RTSP stream")
                
                # Try to get a frame
                ret, frame = False, None
                for _ in range(5):  # Try multiple times
                    ret, frame = cap.read()
                    if ret:
                        break
                
                cap.release()
                
                if ret and frame is not None:
                    # Convert to PIL Image
                    image_pil = Image.fromarray(cv2.cvtColor(frame, cv2.COLOR_BGR2RGB))
                    image_pil.thumbnail(thumbnail_size, Image.LANCZOS)
                    img_byte_arr = io.BytesIO()
                    image_pil.save(img_byte_arr, format='PNG')
                    image_data = img_byte_arr.getvalue()
                else:
                    raise ValueError("Failed to capture frame from RTSP")
                    
            except Exception as e:
                print(f"Error fetching thumbnail from RTSP ({rtsp_url}): {e}")
                image_data = None

        if image_data:
            try:
                img = Image.open(io.BytesIO(image_data))
                img.thumbnail(thumbnail_size, Image.LANCZOS)
                photo = ImageTk.PhotoImage(img)
                self.hover_image_cache[image_key] = photo
                self.root.after(0, lambda: self._update_hover_window(thread_id, photo))
            except Exception as e:
                print(f"Error processing image: {e}")
                self.root.after(0, lambda: self._update_hover_window_error(thread_id, "Error loading image"))
        else:
            self.root.after(0, lambda: self._update_hover_window_error(thread_id, "No image available"))

    def _update_hover_window(self, thread_id, photo):
        """Updates the hover window with the thumbnail."""
        if thread_id != self.last_hover_thread_id or not self.hover_window:
            return

        if self.hover_image_label:
            self.hover_image_label.destroy()
        
        self.current_hover_image_tk = photo
        image_label = ttk.Label(self.hover_window, image=photo)
        image_label.pack()
        self.hover_window.update_idletasks()

    def _update_hover_window_error(self, thread_id, message):
        """Updates the hover window with an error message."""
        if thread_id != self.last_hover_thread_id or not self.hover_window:
            return

        if self.hover_image_label:
            self.hover_image_label.config(text=message)
        else:
            error_label = ttk.Label(self.hover_window, text=message)
            error_label.pack(padx=5, pady=5)
        self.hover_window.update_idletasks()

    def save_project(self):
        """Saves current device data to a JSON file."""
        file_path = filedialog.asksaveasfilename(
            defaultextension=".json",
            filetypes=[("Project Files", "*.json"), ("All Files", "*.*")],
            title="Save Project"
        )
        if not file_path:
            return

        results_to_save = []
        for item_id in self.tree.get_children():
            values = self.tree.item(item_id, 'values')
            device_data = {
                "IP": values[0],
                "MAC": values[1],
                "Credential Set": values[2],
                "Open Ports": values[3],
                "RTSP URL": values[4],
                "Camera Status": values[5],
                "Snapshot": values[6]
            }
            results_to_save.append(device_data)

        try:
            with open(file_path, 'w') as f:
                json.dump(results_to_save, f, indent=4)
            self.show_floating_message(f"Project saved to {os.path.basename(file_path)}")
        except Exception as e:
            self.show_floating_message(f"Error saving project: {e}")

    def load_project(self):
        """Loads device data from a JSON file into the tree."""
        file_path = filedialog.askopenfilename(
            filetypes=[("Project Files", "*.json"), ("All Files", "*.*")],
            title="Load Project"
        )
        if not file_path:
            return

        try:
            with open(file_path, 'r') as f:
                loaded_data = json.load(f)

            self.tree.delete(*self.tree.get_children())
            self.found_count = 0

            for device_data in loaded_data:
                self.add_device_to_tree(device_data) 
            
            self.update_status(f"Project loaded. {self.found_count} devices.")
        except Exception as e:
            self.show_floating_message(f"Error loading project: {e}")

if __name__ == "__main__":
    root = ThemedTk(theme="arc")
    app = RealTimeNetworkScanner(root)
    root.mainloop()
